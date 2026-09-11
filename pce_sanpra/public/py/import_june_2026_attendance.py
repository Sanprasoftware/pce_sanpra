from datetime import date, timedelta
from decimal import Decimal

import frappe
from frappe.utils import cint
from openpyxl import load_workbook


START_DATE = date(2026, 6, 1)
END_DATE = date(2026, 6, 30)
WORKBOOK_NAME = "pec attendance.xlsx"


def import_attendance(dry_run=False):
	"""Import June 2026 attendance totals, excluding every Tuesday."""
	dry_run = cint(dry_run)
	working_dates = _get_working_dates()
	attendance_rows = _get_attendance_rows()
	employees = _match_employees(attendance_rows)
	planned_records, capped_employees = _build_plan(
		attendance_rows, employees, working_dates
	)

	if dry_run:
		return _get_summary(
			attendance_rows,
			working_dates,
			planned_records,
			capped_employees,
		)

	created = 0
	skipped = 0

	for record in planned_records:
		existing = frappe.db.get_value(
			"Attendance",
			{
				"employee": record["employee"],
				"attendance_date": record["attendance_date"],
				"docstatus": ["<", 2],
			},
			["name", "status", "half_day_status", "docstatus"],
			as_dict=True,
		)

		if existing:
			_validate_existing_attendance(existing, record)
			skipped += 1
			continue

		doc = frappe.get_doc(
			{
				"doctype": "Attendance",
				"employee": record["employee"],
				"attendance_date": record["attendance_date"],
				"status": record["status"],
				"half_day_status": record["half_day_status"],
			}
		)
		doc.insert(ignore_permissions=True)
		doc.submit()
		created += 1

	frappe.db.commit()

	return {
		**_get_summary(
			attendance_rows,
			working_dates,
			planned_records,
			capped_employees,
		),
		"created": created,
		"skipped_existing": skipped,
	}


def _get_working_dates():
	working_dates = []
	current_date = START_DATE

	while current_date <= END_DATE:
		if current_date.weekday() != 1:  # Tuesday
			working_dates.append(current_date)
		current_date += timedelta(days=1)

	return working_dates


def _get_attendance_rows():
	workbook_path = frappe.get_app_path(
		"pce_sanpra", "public", "docs", WORKBOOK_NAME
	)
	workbook = load_workbook(workbook_path, read_only=True, data_only=True)
	worksheet = workbook.active
	rows = []

	for row_number, (employee_name, attendance_days) in enumerate(
		worksheet.iter_rows(min_row=2, max_col=2, values_only=True), start=2
	):
		if not employee_name and attendance_days is None:
			continue

		if not employee_name or attendance_days is None:
			frappe.throw(f"Incomplete attendance data in workbook row {row_number}")

		attendance_days = Decimal(str(attendance_days))
		if attendance_days < 0 or attendance_days % Decimal("0.5"):
			frappe.throw(
				f"Attendance days must be a non-negative multiple of 0.5 "
				f"in workbook row {row_number}"
			)

		rows.append(
			{
				"row_number": row_number,
				"employee_name": str(employee_name).strip(),
				"attendance_days": attendance_days,
			}
		)

	workbook.close()
	return rows


def _match_employees(attendance_rows):
	employees_by_name = {}
	for employee in frappe.get_all(
		"Employee",
		filters={"status": "Active"},
		fields=["name", "employee_name"],
	):
		key = _normalize_name(employee.employee_name)
		employees_by_name.setdefault(key, []).append(employee)

	matched = {}
	for row in attendance_rows:
		key = _normalize_name(row["employee_name"])
		matches = employees_by_name.get(key, [])
		if len(matches) != 1:
			frappe.throw(
				f"Expected one active employee match for workbook row "
				f"{row['row_number']} ({row['employee_name']}), found {len(matches)}"
			)
		matched[row["row_number"]] = matches[0]

	return matched


def _normalize_name(employee_name):
	return " ".join(str(employee_name).split()).casefold()


def _build_plan(attendance_rows, employees, working_dates):
	planned_records = []
	capped_employees = []
	maximum_days = Decimal(len(working_dates))

	for row in attendance_rows:
		target_days = min(row["attendance_days"], maximum_days)
		if target_days != row["attendance_days"]:
			capped_employees.append(
				{
					"employee": employees[row["row_number"]].name,
					"employee_name": row["employee_name"],
					"workbook_days": float(row["attendance_days"]),
					"imported_days": float(target_days),
				}
			)

		full_days = int(target_days)
		has_half_day = target_days - full_days == Decimal("0.5")

		for index, attendance_date in enumerate(working_dates):
			if index < full_days:
				status = "Present"
				half_day_status = None
			elif index == full_days and has_half_day:
				status = "Half Day"
				half_day_status = "Absent"
			else:
				status = "Absent"
				half_day_status = None

			planned_records.append(
				{
					"employee": employees[row["row_number"]].name,
					"employee_name": row["employee_name"],
					"attendance_date": attendance_date,
					"status": status,
					"half_day_status": half_day_status,
				}
			)

	return planned_records, capped_employees


def _validate_existing_attendance(existing, planned):
	if existing.docstatus != 1:
		frappe.throw(
			f"Existing Attendance {existing.name} is not submitted for "
			f"{planned['employee']} on {planned['attendance_date']}"
		)

	if existing.status != planned["status"] or (
		(existing.half_day_status or None) != planned["half_day_status"]
	):
		frappe.throw(
			f"Existing Attendance {existing.name} does not match the workbook plan"
		)


def _get_summary(
	attendance_rows, working_dates, planned_records, capped_employees
):
	status_counts = {"Present": 0, "Half Day": 0, "Absent": 0}
	for record in planned_records:
		status_counts[record["status"]] += 1

	return {
		"employees": len(attendance_rows),
		"working_dates": len(working_dates),
		"excluded_tuesdays": 5,
		"planned_records": len(planned_records),
		"status_counts": status_counts,
		"capped_employees": capped_employees,
	}
